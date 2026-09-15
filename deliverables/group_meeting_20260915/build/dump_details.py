# -*- coding: utf-8 -*-
"""只读：导出模板细节 —— 每页每形状的段落结构、run 字体、bodyPr(autofit/wrap)、表格内容。
用法：python3 dump_details.py template_source.pptx > template_details.txt
"""
import sys
from lxml import etree
from pptx import Presentation
from pptx.util import Emu

A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
EMU = 914400.0


def i(v):
    return None if v is None else round(v / EMU, 3)


def rpr_summary(r):
    rPr = r._r.find('{%s}rPr' % A)
    if rPr is None:
        return '<no rPr>'
    out = []
    for tag in ('latin', 'ea', 'cs'):
        e = rPr.find('{%s}%s' % (A, tag))
        if e is not None:
            out.append('%s=%s' % (tag, e.get('typeface')))
    for tag in ('sz', 'b', 'i'):
        e = rPr.find('{%s}%s' % (A, tag))
        if e is not None:
            out.append('%s=%s' % (tag, e.get('val')))
    sf = rPr.find('{%s}solidFill' % A)
    if sf is not None:
        c = sf.find('{%s}srgbClr' % A)
        if c is not None:
            out.append('color=%s' % c.get('val'))
        s = sf.find('{%s}schemeClr' % A)
        if s is not None:
            out.append('color=scheme:%s' % s.get('val'))
    return ','.join(out)


def body_summary(tf):
    bp = tf._txBody.find('{%s}bodyPr' % A)
    out = {}
    if bp is not None:
        for k in ('wrap', 'anchor', 'lIns', 'tIns', 'rIns', 'bIns', 'rot', 'vert'):
            if bp.get(k) is not None:
                out[k] = bp.get(k)
        for tag in ('spAutoFit', 'normAutofit', 'noAutofit'):
            if bp.find('{%s}%s' % (A, tag)) is not None:
                out['autofit'] = tag
    return out


def dump_shape(sh, prefix, out, depth=0):
    pad = '  ' * depth
    if sh.shape_type == 6 and sh.__class__.__name__ == 'GroupShape':
        out.append('%s%s [GROUP]' % (pad, prefix + sh.name))
        for s in sh.shapes:
            dump_shape(s, prefix, out, depth + 1)
        return
    if getattr(sh, 'has_table', False) and sh.has_table:
        t = sh.table
        out.append('%s%s [TABLE %dx%d] L=%s T=%s W=%s H=%s' % (
            pad, sh.name, len(t.rows), len(t.columns), i(sh.left), i(sh.top), i(sh.width), i(sh.height)))
        for ri, row in enumerate(t.rows):
            out.append('%s  row%d h=%s' % (pad, ri, i(row.height)))
            for ci, cell in enumerate(row.cells):
                txt = cell.text.replace('\n', '\\n')
                out.append('%s    c%d len=%d | %s' % (pad, ci, len(txt), txt[:90]))
        return
    if not sh.has_text_frame:
        out.append('%s%s | L=%s T=%s W=%s H=%s | (no text frame)' % (
            pad, sh.name, i(sh.left), i(sh.top), i(sh.width), i(sh.height)))
        return
    tf = sh.text_frame
    out.append('%s%s | L=%s T=%s W=%s H=%s | len=%d | bodyPr=%s' % (
        pad, sh.name, i(sh.left), i(sh.top), i(sh.width), i(sh.height), len(tf.text), body_summary(tf)))
    for pi, p in enumerate(tf.paragraphs):
        runs = ' ;; '.join('%s{%s}' % (r.text, rpr_summary(r)) for r in p.runs)
        pPr = p._p.find('{%s}pPr' % A)
        algn = pPr.get('algn') if pPr is not None else None
        lvl = pPr.get('lvl') if pPr is not None else None
        out.append('%s  p%d len=%d algn=%s lvl=%s | %s' % (
            pad, pi, len(p.text), algn, lvl, runs))


def main():
    prs = Presentation(sys.argv[1])
    out = []
    for idx, slide in enumerate(prs.slides, 1):
        out.append('')
        out.append('#' * 90)
        out.append('SLIDE %d layout=%s' % (idx, slide.slide_layout.name))
        out.append('#' * 90)
        for sh in slide.shapes:
            dump_shape(sh, '', out)
        if slide.has_notes_slide:
            nt = slide.notes_slide.notes_text_frame
            out.append('  [NOTES] bodyPr=%s' % body_summary(nt))
            for pi, p in enumerate(nt.paragraphs):
                out.append('  n%d len=%d | %s' % (pi, len(p.text), ' ;; '.join(
                    '%s{%s}' % (r.text, rpr_summary(r)) for r in p.runs)))
    sys.stdout.write('\n'.join(out) + '\n')


if __name__ == '__main__':
    main()
